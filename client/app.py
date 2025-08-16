import os, time, requests
from pathlib import Path
import gradio as gr
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfReader

APP_DIR = Path(__file__).parent
DATA_DIR = APP_DIR / "data"
SECRETS_DIR = APP_DIR / ".secrets"
for p in (DATA_DIR, SECRETS_DIR):
    p.mkdir(parents=True, exist_ok=True)

BACKEND_FILE = SECRETS_DIR / "backend_url.txt"
LICENSE_FILE = SECRETS_DIR / "license_key.txt"

def save_backend(url: str):
    if not url.startswith("http"):
        raise gr.Error("Enter a full URL, e.g., http://localhost:8000")
    BACKEND_FILE.write_text(url.strip(), encoding="utf-8")
    return "✅ Saved"

def save_license(key: str):
    if len(key.strip()) < 8:
        raise gr.Error("License looks too short.")
    LICENSE_FILE.write_text(key.strip(), encoding="utf-8")
    return "✅ Saved"

def get_backend():
    return (BACKEND_FILE.read_text(encoding="utf-8").strip() if BACKEND_FILE.exists() else "http://localhost:8000")

def get_license():
    return (LICENSE_FILE.read_text(encoding="utf-8").strip() if LICENSE_FILE.exists() else "")

def check_license():
    url = get_backend() + "/api/license/verify"
    key = get_license()
    if not key: return "❌ No license saved"
    try:
        resp = requests.post(url, json={"key": key}, timeout=15)
        resp.raise_for_status()
        data = resp.json()
        return f"✅ Plan: {data['plan']} | Remaining: {data['quota_remaining']}"
    except Exception as e:
        return f"❌ {e}"

def upload_file(files):
    if not files: return "No file selected."
    url = get_backend() + "/api/upload"
    key = get_license()
    if not key: raise gr.Error("Save a license key first.")
    total_chunks = 0
    for f in files:
        with open(f.name, "rb") as fh:
            files_data = {"file": (Path(f.name).name, fh, "application/octet-stream")}
            data = {"key": key}
            r = requests.post(url, data=data, files=files_data, timeout=120)
            if r.status_code != 200:
                return f"❌ Upload failed for {f.name}: {r.text}"
            total_chunks += r.json().get("chunks_added", 0)
    return f"✅ Uploaded. {total_chunks} chunks added."

def chat(msg, mode):
    url = get_backend() + "/api/chat"
    key = get_license()
    if not key: raise gr.Error("Save a license key first.")
    try:
        r = requests.post(url, json={"key": key, "message": msg, "mode": mode}, timeout=120)
        r.raise_for_status()
        data = r.json()
        return data["answer"], "Citations: " + ", ".join(data.get("citations", []))
    except Exception as e:
        return f"❌ {e}", ""

def export_docx(content: str):
    out = DATA_DIR / f"finding_{time.strftime('%Y%m%d-%H%M%S')}.docx"
    doc = DocxDocument()
    for line in content.splitlines() or [""]:
        doc.add_paragraph(line)
    doc.save(str(out))
    return str(out)

def export_xlsx(content: str):
    out = DATA_DIR / f"export_{time.strftime('%Y%m%d-%H%M%S')}.xlsx"
    wb = Workbook()
    ws1 = wb.active; ws1.title = "Findings"
    headers = ["Control No","Control Name","Question","Example how to demonstrate question","Implementation","Evidence","C","NC","OFI","NA","Recommendations"]
    ws1.append(headers)
    ws2 = wb.create_sheet("Output")
    ws2.append(["Output"])
    for line in content.splitlines() or [""]:
        ws2.append([line])
    wb.save(str(out))
    return str(out)

with gr.Blocks(title="Bitseclab Copilot") as demo:
    gr.Markdown("# Bitseclab Copilot (License-gated)")
    with gr.Row():
        with gr.Column():
            backend = gr.Textbox(label="Backend URL", value=(get_backend()), placeholder="http://localhost:8000")
            save_backend_btn = gr.Button("Save Backend")
            backend_status = gr.Markdown("")
            lic = gr.Textbox(label="License Key", value=(get_license()))
            save_lic_btn = gr.Button("Save License")
            lic_status = gr.Markdown("")
            verify_btn = gr.Button("Verify License")
            verify_status = gr.Markdown("")
        with gr.Column(scale=2):
            msg = gr.Textbox(label="Your message", lines=4)
            mode = gr.Radio(["answer","draft","critique"], value="answer", label="Mode")
            send = gr.Button("Send", variant="primary")
            out = gr.Markdown(label="Output")
            cites = gr.Markdown(label="Citations")
            with gr.Row():
                btn_docx = gr.Button("Export to Word (.docx)")
                btn_xlsx = gr.Button("Export to Excel (.xlsx)")
            file_docx = gr.File(label="Download DOCX", interactive=False)
            file_xlsx = gr.File(label="Download XLSX", interactive=False)
            up = gr.File(label="Upload files (PDF/DOCX/TXT/MD)", file_count="multiple")
            ingest_btn = gr.Button("Upload & Ingest")

    save_backend_btn.click(save_backend, inputs=[backend], outputs=[backend_status])
    save_lic_btn.click(save_license, inputs=[lic], outputs=[lic_status])
    verify_btn.click(lambda: check_license(), outputs=[verify_status])
    send.click(chat, inputs=[msg, mode], outputs=[out, cites])
    ingest_btn.click(upload_file, inputs=[up], outputs=[verify_status])
    btn_docx.click(lambda t: export_docx(t) if t else "", inputs=[out], outputs=[file_docx])
    btn_xlsx.click(lambda t: export_xlsx(t) if t else "", inputs=[out], outputs=[file_xlsx])

if __name__ == "__main__":
    # show_api=False avoids 3.13 schema issue
    demo.launch(server_name="0.0.0.0", server_port=7860, share=True, show_api=False, inbrowser=True)